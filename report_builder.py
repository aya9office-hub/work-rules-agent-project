from __future__ import annotations
from datetime import date
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def build_report(findings: list[dict], input_files: list[str], out_path: Path) -> Path:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Yu Gothic"
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_heading("就業規則・附属規程 構造チェック報告書", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"作成日: {date.today().isoformat()}")

    doc.add_heading("1. チェックの目的", level=1)
    doc.add_paragraph("本報告書は、法違反の断定を目的とせず、就業規則および附属規程を横断して、規程間の矛盾、参照関係、用語の不統一、運用上の不明確さなど、構造上の問題点を確認するものです。")

    doc.add_heading("2. 対象ファイル", level=1)
    for f in input_files:
        doc.add_paragraph(f, style="List Bullet")

    doc.add_heading("3. 総評", level=1)
    if findings:
        doc.add_paragraph(f"構造上の確認事項が {len(findings)} 件検出されました。重要度の高い項目から順に内容を確認してください。")
    else:
        doc.add_paragraph("本チェック範囲では、明確な構造上の問題点は検出されませんでした。ただし、最終判断は原規程と運用実態を確認したうえで行ってください。")

    doc.add_heading("4. 検出事項", level=1)
    if findings:
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        headers = ["No", "分類", "重要度", "箇所", "問題点", "確認・修正案"]
        for i, h in enumerate(headers):
            run = table.rows[0].cells[i].paragraphs[0].add_run(h)
            run.bold = True
        for idx, f in enumerate(findings, 1):
            cells = table.add_row().cells
            cells[0].text = str(idx)
            cells[1].text = f.get("category", "")
            cells[2].text = f.get("severity", "")
            cells[3].text = f.get("location", "")
            cells[4].text = f"{f.get('title','')}\n{f.get('issue','')}"
            cells[5].text = f.get("recommendation", "")
    else:
        doc.add_paragraph("検出事項なし。")

    doc.add_heading("5. 注意事項", level=1)
    doc.add_paragraph("本報告書はAIによる一次チェック結果です。法的評価、行政解釈、最新法改正への適合性を保証するものではありません。実務上の判断は、専門家による最終確認を前提としてください。")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
